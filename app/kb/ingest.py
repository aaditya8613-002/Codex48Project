"""Read guideline files (PDF, DOCX, TXT, MD), split them into passages, store them."""
import io
import re

from sqlalchemy.orm import Session

from ..ai import gemini
from ..models import BrandDocument, KbChunk

MAX_CHARS = 900
OVERLAP = 120


def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        import pdfplumber

        with pdfplumber.open(io.BytesIO(data)) as pdf:
            return "\n\n".join((page.extract_text() or "") for page in pdf.pages)
    if name.endswith(".docx"):
        import docx

        d = docx.Document(io.BytesIO(data))
        lines = []
        for p in d.paragraphs:
            if not p.text.strip():
                lines.append("")
            elif p.style is not None and p.style.name.lower().startswith("heading"):
                lines.append("# " + p.text.strip())
            else:
                lines.append(p.text.strip())
        return "\n".join(lines)
    if name.endswith((".txt", ".md", ".markdown")):
        return data.decode("utf-8", errors="replace")
    raise ValueError("Unsupported file type. Upload PDF, DOCX, TXT or MD.")


def chunk_text(text: str) -> list[tuple[str, str]]:
    """Split into (heading, passage). Keeps paragraphs whole where possible."""
    text = text.replace("\r\n", "\n")
    heading = ""
    chunks: list[tuple[str, str]] = []
    buf = ""

    def flush():
        nonlocal buf
        if buf.strip():
            chunks.append((heading, buf.strip()))
        buf = ""

    for block in re.split(r"\n\s*\n", text):
        block = block.strip()
        if not block:
            continue
        first, _, rest = block.partition("\n")
        if first.startswith("#"):
            flush()
            heading = first.lstrip("#").strip()
            block = rest.strip()
            if not block:
                continue
        if len(buf) + len(block) + 2 <= MAX_CHARS:
            buf = f"{buf}\n\n{block}" if buf else block
            continue
        flush()
        if len(block) <= MAX_CHARS:
            buf = block
            continue
        start = 0  # a single very long paragraph: sliding window
        while start < len(block):
            piece = block[start : start + MAX_CHARS]
            chunks.append((heading, piece.strip()))
            start += MAX_CHARS - OVERLAP
    flush()
    return chunks


def add_document(
    session: Session,
    org_id: int,
    title: str,
    text: str,
    filename: str = "",
    kind: str = "guideline",
    embed: bool = True,
) -> BrandDocument:
    doc = BrandDocument(org_id=org_id, title=title, filename=filename, kind=kind)
    pieces = chunk_text(text)
    if not pieces:
        raise ValueError("No readable text found in the document.")
    vectors: list[list[float]] | list[None] = [None] * len(pieces)
    if embed and gemini.is_configured():
        try:
            vectors = gemini.embed_texts([f"{h}\n{t}" if h else t for h, t in pieces])
        except gemini.AIUnavailable:
            vectors = [None] * len(pieces)  # can be filled later by reembed_missing()
    for i, ((heading, passage), vec) in enumerate(zip(pieces, vectors)):
        doc.chunks.append(KbChunk(org_id=org_id, position=i, heading=heading, text=passage, embedding=vec))
    session.add(doc)
    session.commit()
    return doc


def reembed_missing(session: Session, org_id: int) -> int:
    """Embed any passage that was saved without a vector (for example after an API outage)."""
    rows = session.query(KbChunk).filter(KbChunk.org_id == org_id, KbChunk.embedding.is_(None)).all()
    if not rows:
        return 0
    vectors = gemini.embed_texts([f"{r.heading}\n{r.text}" if r.heading else r.text for r in rows])
    for row, vec in zip(rows, vectors):
        row.embedding = vec
    session.commit()
    return len(rows)
